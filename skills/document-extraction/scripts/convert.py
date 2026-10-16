#!/usr/bin/env python3
# /// script
# requires-python = ">=3.11"
# dependencies = [
#   "pdfplumber>=0.11",
#   "python-docx>=1.1",
#   "openpyxl>=3.1",
#   "python-pptx>=1.0",
# ]
# ///
"""Convert a PDF/DOCX/XLSX/PPTX document to a JSON intermediate representation.

Writes <out>/ir.json and prints a summary (pages/sheets/slides, tables, warnings)
to stdout. Every value in the IR is coordinate-addressable (page/table/row for
PDFs, sheet/cell for spreadsheets, slide/shape for decks) so extracted values can
be traced back to a location in the source document.

IR shape (top level): {"source", "kind", "warnings", ...kind-specific}
  pdf:  {"pages":  [{"page": 1, "text": "...", "no_text_layer": false,
                     "tables": [{"index": 1, "rows": [[cell, ...], ...]}]}]}
  docx: {"blocks": [{"type": "heading", "level": 1, "text": "..."} |
                    {"type": "paragraph", "text": "..."} |
                    {"type": "table", "index": 1, "rows": [[cell, ...], ...]}]}
  xlsx: {"sheets": [{"name": "...", "truncated": false,
                     "rows": [{"row": 1, "cells": {"A": v, "C": v}}]}]}
        (empty cells omitted; values keep their types, dates as ISO strings)
  pptx: {"slides": [{"slide": 1, "texts": ["..."],
                     "tables": [{"index": 1, "rows": [[cell, ...], ...]}],
                     "notes": "..." | null}]}
"""
from __future__ import annotations

import argparse
import json
import sys
from datetime import date, datetime, time
from pathlib import Path


def clean_cell(v: object) -> object:
    """Normalize a raw cell for JSON: trim strings, ISO-format dates."""
    if isinstance(v, str):
        return v.strip()
    if isinstance(v, datetime):
        return v.date().isoformat() if (v.hour, v.minute, v.second) == (0, 0, 0) else v.isoformat()
    if isinstance(v, (date, time)):
        return v.isoformat()
    return v


def clean_rows(rows: list[list[object]]) -> list[list[object]]:
    return [[clean_cell(c) for c in r] for r in rows]


def convert_pdf(path: Path, warnings: list[str]) -> dict:
    import pdfplumber

    pages = []
    n_tables = 0
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages, 1):
            text = (page.extract_text() or "").strip()
            tables = page.extract_tables() or []
            no_text = not text and not tables
            if no_text:
                warnings.append(f"NO_TEXT_LAYER: page {i} has no extractable text (scanned image?)")
            n_tables += len(tables)
            pages.append({
                "page": i,
                "no_text_layer": no_text,
                "text": text,
                "tables": [{"index": t + 1, "rows": clean_rows(tbl)} for t, tbl in enumerate(tables)],
            })
    return {"kind": "pdf", "pages": pages, "stats": {"pages": len(pages), "tables": n_tables}}


def convert_docx(path: Path, warnings: list[str]) -> dict:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = docx.Document(str(path))
    blocks: list[dict] = []
    n_tables = 0
    n_paras = 0

    # Iterate body elements in document order (paragraphs and tables interleaved).
    for child in doc.element.body.iterchildren():
        if child.tag.endswith("}p"):
            p = Paragraph(child, doc)
            text = p.text.strip()
            if not text:
                continue
            n_paras += 1
            style = (p.style.name or "").lower() if p.style is not None else ""
            if style.startswith("title"):
                blocks.append({"type": "heading", "level": 0, "text": text})
            elif style.startswith("heading"):
                try:
                    level = int(style.split()[-1])
                except ValueError:
                    level = 2
                blocks.append({"type": "heading", "level": level, "text": text})
            else:
                block: dict = {"type": "paragraph", "text": text}
                if style.startswith("list"):
                    block["style"] = style
                blocks.append(block)
        elif child.tag.endswith("}tbl"):
            n_tables += 1
            table = Table(child, doc)
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            blocks.append({"type": "table", "index": n_tables, "rows": rows})
    return {"kind": "docx", "blocks": blocks, "stats": {"paragraphs": n_paras, "tables": n_tables}}


def convert_xlsx(
    path: Path, warnings: list[str], max_rows: int, max_cols: int, only_sheet: str | None
) -> dict:
    import openpyxl
    from openpyxl.utils import get_column_letter

    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    sheets = []
    for ws in wb.worksheets:
        if only_sheet and ws.title != only_sheet:
            continue
        rows = list(ws.iter_rows(values_only=True))
        while rows and all(c is None for c in rows[-1]):
            rows.pop()
        truncated = False
        if len(rows) > max_rows:
            warnings.append(
                f"TRUNCATED: sheet '{ws.title}' has {len(rows)} rows; showing first {max_rows} "
                f"(re-run with --max-rows to see more)"
            )
            rows, truncated = rows[:max_rows], True
        ncols = max((len(r) for r in rows), default=0)
        if ncols > max_cols:
            warnings.append(f"TRUNCATED: sheet '{ws.title}' has {ncols} columns; showing first {max_cols}")
            ncols, truncated = max_cols, True
        json_rows = []
        for i, r in enumerate(rows, 1):
            cells = {
                get_column_letter(c + 1): clean_cell(r[c])
                for c in range(min(ncols, len(r)))
                if r[c] is not None
            }
            if cells:
                json_rows.append({"row": i, "cells": cells})
        sheets.append({"name": ws.title, "truncated": truncated, "rows": json_rows})
    wb.close()
    return {"kind": "xlsx", "sheets": sheets, "stats": {"sheets": len(sheets)}}


def convert_pptx(path: Path, warnings: list[str]) -> dict:
    from pptx import Presentation

    prs = Presentation(str(path))
    slides = []
    n_tables = 0
    for i, slide in enumerate(prs.slides, 1):
        texts: list[str] = []
        tables: list[dict] = []
        for shape in slide.shapes:
            if shape.has_table:
                n_tables += 1
                rows = [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
                tables.append({"index": n_tables, "rows": rows})
            elif shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    texts.append(text)
        notes = None
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip() or None
        slides.append({"slide": i, "texts": texts, "tables": tables, "notes": notes})
    return {"kind": "pptx", "slides": slides, "stats": {"slides": len(slides), "tables": n_tables}}


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("document", type=Path)
    ap.add_argument("--out", type=Path, default=None, help="output directory (default: alongside doc)")
    ap.add_argument("--max-rows", type=int, default=200, help="xlsx: max rows rendered per sheet")
    ap.add_argument("--max-cols", type=int, default=40, help="xlsx: max columns rendered per sheet")
    ap.add_argument("--sheet", default=None, help="xlsx: render only this sheet")
    args = ap.parse_args()

    doc: Path = args.document
    if not doc.exists():
        print(f"error: {doc} not found", file=sys.stderr)
        return 2
    out_dir = args.out or doc.parent
    out_dir.mkdir(parents=True, exist_ok=True)

    warnings: list[str] = []
    ext = doc.suffix.lower()
    converters = {
        ".pdf": lambda: convert_pdf(doc, warnings),
        ".docx": lambda: convert_docx(doc, warnings),
        ".xlsx": lambda: convert_xlsx(doc, warnings, args.max_rows, args.max_cols, args.sheet),
        ".pptx": lambda: convert_pptx(doc, warnings),
    }
    if ext not in converters:
        print(f"error: unsupported extension {ext} (supported: {', '.join(converters)})", file=sys.stderr)
        return 2

    ir = converters[ext]()
    ir = {"source": doc.name, "kind": ir.pop("kind"), "warnings": warnings, **ir}
    stats = ir.pop("stats")
    ir_path = out_dir / "ir.json"
    ir_path.write_text(json.dumps(ir, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")

    print(json.dumps({"ir": str(ir_path), "warnings": warnings, "kind": ir["kind"], **stats}, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
